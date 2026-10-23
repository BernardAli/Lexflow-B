"""Document parser. No outbound calls; OCR uses an optional local Tesseract installation."""
import email
import io
from email import policy
from html.parser import HTMLParser
from pathlib import Path
from fastapi import HTTPException
from PIL import Image
from .settings import max_upload_bytes

SUPPORTED = {".txt", ".eml", ".pdf", ".docx", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}
MAX_TEXT = 200_000
MAX_OCR_PAGES = 12


class HTMLText(HTMLParser):
    """Read email text without executing or fetching its HTML content."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "head"}:
            self.hidden += 1
        if tag in {"p", "div", "br", "li", "tr", "h1", "h2"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "head"}:
            self.hidden = max(0, self.hidden - 1)
        if tag in {"p", "div", "li", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def limited_text(text, warnings):
    if len(text) > MAX_TEXT:
        warnings.append(f"Extracted text limited to {MAX_TEXT:,} characters; remaining content was not processed.")
    return text[:MAX_TEXT], warnings


def pdf_text(raw, warnings):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(raw), strict=False)
    if len(reader.pages) > 75:
        raise HTTPException(status_code=413, detail="PDF exceeds 75-page limit")
    parts = [(page.extract_text(extraction_mode="plain") or "") for page in reader.pages]
    # Check each page: a text cover sheet must not suppress OCR of scanned attachments.
    scan_pages = [i for i, text in enumerate(parts) if len(text.strip()) < 30]
    if not scan_pages:
        return limited_text("\n".join(parts), warnings)
    if len(scan_pages) > MAX_OCR_PAGES:
        warnings.append(f"OCR limited to {MAX_OCR_PAGES} PDF pages; {len(scan_pages)-MAX_OCR_PAGES} additional pages need manual inspection.")
    pdf = None
    try:
        import pytesseract
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(raw)
        for i in scan_pages[:MAX_OCR_PAGES]:
            page = bitmap = None
            try:
                page = pdf[i]
                bitmap = page.render(scale=1.8)
                with bitmap.to_pil() as image:
                    recognized = pytesseract.image_to_string(image, lang="deu+eng", timeout=30)
                if recognized.strip():
                    parts[i] = recognized
                    warnings.append(f"Local OCR used for PDF page {i+1}; verify recognized text.")
                else:
                    warnings.append(f"No text recognized on PDF page {i+1}; inspect the original page.")
            except Exception:
                warnings.append(f"OCR unavailable for PDF page {i+1}; inspect the original or paste its text.")
            finally:
                if bitmap is not None:
                    bitmap.close()
                if page is not None:
                    page.close()
    except Exception:
        warnings.append("PDF OCR unavailable. Install Tesseract (German/English) or paste text; pages with little or no text need inspection.")
    finally:
        if pdf is not None:
            pdf.close()
    return limited_text("\n".join(parts), warnings)


def email_text(raw, warnings, depth):
    if depth >= 3:
        warnings.append("Nested email depth limit reached; inspect remaining attachments manually.")
        return "", warnings
    message = email.message_from_bytes(raw, policy=policy.default)
    headers = "\n".join(f"{key}: {message.get(key, '')}" for key in ("From", "To", "Subject"))
    blocks = [headers]
    body = message.get_body(preferencelist=("plain", "html"))
    if body is not None:
        content = body.get_content()
        if body.get_content_type() == "text/html":
            parser = HTMLText()
            parser.feed(content)
            content = "".join(parser.parts)
        blocks.append(content)
    attachments = list(message.iter_attachments())
    if len(attachments) > 20:
        warnings.append("Only the first 20 email attachments were processed; inspect the remaining attachments.")
    for part in attachments[:20]:
        filename = Path(part.get_filename() or "unnamed-attachment").name
        if part.get_content_type() == "message/rfc822":
            nested = part.get_payload()
            payload = nested[0].as_bytes() if isinstance(nested, list) and nested else b""
            if not filename.lower().endswith(".eml"):
                filename += ".eml"
        else:
            payload = part.get_payload(decode=True) or b""
        if Path(filename).suffix.lower() not in SUPPORTED:
            warnings.append(f"Attachment {filename} was not processed: unsupported type.")
            continue
        try:
            text, notes = parse_document(filename, payload, _depth=depth+1)
            # A labeled boundary also keeps candidate evidence attributable to the attachment.
            blocks.append(f"\n[Attachment: {filename}]\n{text}")
            warnings.extend(f"Attachment {filename}: {note}" for note in notes)
        except HTTPException as exc:
            # Malformed supported attachments fail the intake, so no incomplete batch is stored.
            raise HTTPException(status_code=exc.status_code, detail=f"Attachment {filename}: {exc.detail}") from exc
    return limited_text("\n".join(blocks), warnings)


def parse_document(filename: str, raw: bytes, *, _depth: int = 0) -> tuple[str, list[str]]:
    if len(raw) > max_upload_bytes():
        raise HTTPException(status_code=413, detail="File exceeds maximum upload size")
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise HTTPException(status_code=415, detail=f"Unsupported file extension: {suffix}")
    warnings = []
    try:
        if suffix == ".txt":
            return limited_text(raw.decode("utf-8-sig", errors="replace"), warnings)
        if suffix == ".pdf":
            return pdf_text(raw, warnings)
        if suffix == ".docx":
            from docx import Document
            document = Document(io.BytesIO(raw))
            blocks = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    values = [cell.text.strip() for cell in row.cells]
                    if len(values) == 2 and values[0] and values[1]:
                        blocks.append(f"{values[0]}: {values[1]}")
                    else:
                        blocks.append("\t".join(values))
            return limited_text("\n".join(blocks), warnings)
        if suffix == ".eml":
            return email_text(raw, warnings, _depth)
        if suffix in (".png", ".jpg", ".jpeg", ".tif", ".tiff"):
            with Image.open(io.BytesIO(raw)) as image:
                frames = getattr(image, "n_frames", 1)
                if frames > MAX_OCR_PAGES:
                    warnings.append(f"OCR limited to {MAX_OCR_PAGES} image frames; {frames-MAX_OCR_PAGES} additional frames need inspection.")
                parts = []
                for i in range(min(frames, MAX_OCR_PAGES)):
                    image.seek(i)
                    try:
                        import pytesseract
                        text = pytesseract.image_to_string(image, lang="deu+eng", timeout=30)
                        parts.append(text)
                        if not text.strip():
                            warnings.append(f"No text recognized on image frame {i+1}; inspect the original.")
                    except Exception:
                        warnings.append(f"OCR unavailable for image frame {i+1}. Install Tesseract (German/English) or paste its text.")
                return limited_text("\n".join(parts), warnings)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not parse document: {type(exc).__name__}") from exc
    raise HTTPException(status_code=415, detail="Unsupported file")
